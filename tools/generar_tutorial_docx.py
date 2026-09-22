# -*- coding: utf-8 -*-
"""Genera el tutorial del Panel MT3 en Word (.docx) a partir de TUTORIAL.md.

Uso:  python tools/generar_tutorial_docx.py
Salida: Tutorial_Panel_MT3.docx (raíz del proyecto)
Luego exportar a PDF con el script COM de Word (skill documentos-profesionales-office).
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE = Path(__file__).resolve().parent.parent
MD = BASE / "TUTORIAL.md"
OUT = BASE / "Tutorial_Panel_MT3.docx"

AZUL = RGBColor(0x1F, 0x4E, 0x78)
AZUL_CLARO = "D9E2F3"
GRIS = RGBColor(0x64, 0x74, 0x8B)
FUENTE = "Calibri"


# --------------------------------------------------------------------------- #
# Utilidades de bajo nivel
# --------------------------------------------------------------------------- #
def _sombrear(celda, color_hex: str):
    tcPr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    tcPr.append(shd)


def _campo(par, instr: str):
    """Inserta un campo de Word (p. ej. TOC, PAGE, NUMPAGES)."""
    r1 = par.add_run()._r
    fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "begin"); r1.append(fc)
    r2 = par.add_run()._r
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr; r2.append(it)
    r3 = par.add_run()._r
    fs = OxmlElement("w:fldChar"); fs.set(qn("w:fldCharType"), "separate"); r3.append(fs)
    r4 = par.add_run("Actualiza el indice (F9)")._r
    r5 = par.add_run()._r
    fe = OxmlElement("w:fldChar"); fe.set(qn("w:fldCharType"), "end"); r5.append(fe)


def _runs_inline(par, texto: str, base_size=10.5, color=None):
    """Escribe texto con **negritas** y `codigo` en el párrafo."""
    partes = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", texto)
    for parte in partes:
        if not parte:
            continue
        if parte.startswith("**") and parte.endswith("**") and len(parte) > 4:
            r = par.add_run(parte[2:-2]); r.bold = True
        elif parte.startswith("`") and parte.endswith("`") and len(parte) > 2:
            r = par.add_run(parte[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(base_size - 0.5)
        else:
            r = par.add_run(parte)
        r.font.size = Pt(base_size)
        r.font.name = FUENTE
        if color is not None:
            r.font.color.rgb = color


# --------------------------------------------------------------------------- #
# Documento
# --------------------------------------------------------------------------- #
def nuevo_doc() -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Cm(2.2)
    sec.left_margin = sec.right_margin = Cm(2.4)

    st = doc.styles["Normal"]
    st.font.name = FUENTE
    st.font.size = Pt(10.5)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.12

    for nombre, tam, color in (("Heading 1", 15, AZUL), ("Heading 2", 12.5, AZUL),
                               ("Heading 3", 11, RGBColor(0x2E, 0x75, 0xB6))):
        s = doc.styles[nombre]
        s.font.name = "Calibri Light"
        s.font.size = Pt(tam)
        s.font.color.rgb = color
        s.font.bold = True
        s.paragraph_format.space_before = Pt(12 if nombre != "Heading 1" else 16)
        s.paragraph_format.space_after = Pt(5)

    # Encabezado y pie
    hp = sec.header.paragraphs[0]
    hp.text = "Panel MT3 - Tutorial de uso"
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for r in hp.runs:
        r.font.size = Pt(8.5); r.font.color.rgb = GRIS
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run("Pagina "); r.font.size = Pt(8.5); r.font.color.rgb = GRIS
    _campo(fp, "PAGE")
    r2 = fp.add_run(" de "); r2.font.size = Pt(8.5); r2.font.color.rgb = GRIS
    _campo(fp, "NUMPAGES")
    return doc


def portada(doc: Document):
    for _ in range(4):
        doc.add_paragraph()
    t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("Panel MT3"); r.font.size = Pt(30); r.bold = True; r.font.color.rgb = AZUL
    t2 = doc.add_paragraph(); t2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = t2.add_run("Control Mantenimiento Preventivo 3")
    r2.font.size = Pt(16); r2.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)
    t3 = doc.add_paragraph(); t3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r3 = t3.add_run("Tutorial de uso"); r3.font.size = Pt(21); r3.bold = True
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r4 = p.add_run("Guia para usuarios y administradores\nBanco Agrario - COLSOF")
    r4.font.size = Pt(11); r4.font.color.rgb = GRIS
    doc.add_paragraph()
    # Indice
    h = doc.add_heading("Contenido", level=1)
    _campo(doc.add_paragraph(), 'TOC \\o "1-4" \\h \\z \\u')
    doc.add_page_break()


def tabla_md(doc: Document, filas: list):
    n_cols = max(len(f) for f in filas)
    tb = doc.add_table(rows=0, cols=n_cols)
    tb.style = "Table Grid"
    tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, fila in enumerate(filas):
        celdas = tb.add_row().cells
        for j in range(n_cols):
            txt = fila[j] if j < len(fila) else ""
            par = celdas[j].paragraphs[0]
            par.paragraph_format.space_after = Pt(2)
            _runs_inline(par, txt, base_size=9.5)
            if i == 0:
                _sombrear(celdas[j], AZUL_CLARO)
                for run in par.runs:
                    run.bold = True
    doc.add_paragraph()


def convertir(md_texto: str, doc: Document):
    lineas = md_texto.splitlines()
    i = 0
    while i < len(lineas):
        ln = lineas[i].rstrip()
        if ln.startswith("|") and i + 1 < len(lineas) and set(lineas[i + 1].replace("|", "").strip()) <= set("-: "):
            filas = []
            while i < len(lineas) and lineas[i].strip().startswith("|"):
                if not set(lineas[i].replace("|", "").strip()) <= set("-: "):
                    filas.append([c.strip() for c in lineas[i].strip().strip("|").split("|")])
                i += 1
            tabla_md(doc, filas)
            continue
        if ln.startswith("```"):
            i += 1
            codigo = []
            while i < len(lineas) and not lineas[i].startswith("```"):
                codigo.append(lineas[i]); i += 1
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Cm(0.5)
            r = par.add_run("\n".join(codigo))
            r.font.name = "Consolas"; r.font.size = Pt(8.5)
        elif re.match(r"^#{1,4} ", ln):
            nivel = len(ln) - len(ln.lstrip("#"))
            doc.add_heading(ln[nivel + 1:].strip(), level=min(nivel, 4))
        elif re.match(r"^\s*[-*] ", ln):
            par = doc.add_paragraph(style="List Bullet")
            _runs_inline(par, re.sub(r"^\s*[-*] ", "", ln))
        elif re.match(r"^\s*\d+\. ", ln):
            par = doc.add_paragraph(style="List Number")
            _runs_inline(par, re.sub(r"^\s*\d+\. ", "", ln))
        elif ln.startswith(">"):
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Cm(0.6)
            _runs_inline(par, ln.lstrip("> ").strip(), base_size=9.5, color=GRIS)
        elif ln.strip() == "---":
            doc.add_paragraph()
        elif ln.strip():
            par = doc.add_paragraph()
            _runs_inline(par, ln.strip())
        i += 1


def main():
    doc = nuevo_doc()
    portada(doc)
    texto = MD.read_text(encoding="utf-8")
    texto = texto.split("\n", 1)[1] if texto.startswith("# ") else texto  # sin el H1 (va en portada)
    convertir(texto, doc)
    doc.save(OUT)
    print("OK:", OUT)


if __name__ == "__main__":
    main()
