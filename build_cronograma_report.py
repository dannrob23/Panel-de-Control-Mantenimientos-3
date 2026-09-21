# -*- coding: utf-8 -*-
"""
Verificación de cumplimiento / desviación del CRONOGRAMA MT3 (archivo con columna ESTADO)
vs Data_PCTriage.xlsx.

Reglas:
  * Estado de oficina = columna ESTADO del cronograma (Programada / En proceso / Finalizada / Reprogramada).
  * El PROGRESO se mide sobre elementos SI FACTURABLES (columna Facturable='Si' de Data):
        Avance % = MT3 realizados (Fact Si) / Total Fact Si
    Los NO FACTURABLES (COLSOF) se muestran pero no miden el avance.
  * Cumplimiento/desviación:
        Finalizada + pend. Fact Si>0  -> VENCIDA CON PENDIENTES (Fact Si)
        En proceso + sin MT3 Fact Si  -> ATRASADA SIN INICIAR
        Programada + ya con MT3 Si    -> ADELANTADA
        Reprogramada                  -> REPROGRAMADA (revisar)
  * Coherencia: cruce entre ESTADO y las fechas Inicio/Fin del cronograma.
"""
import warnings
warnings.simplefilter("ignore")
import os
from datetime import timezone

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

try:
    import zoneinfo
    TZ_BO = zoneinfo.ZoneInfo("America/Bogota")
except Exception:
    TZ_BO = timezone.utc

def ahora() -> pd.Timestamp:
    return pd.Timestamp.now(tz=TZ_BO).tz_convert(None)

def hoy() -> pd.Timestamp:
    return ahora().normalize()

BASE = os.path.dirname(os.path.abspath(__file__))
CRONO = os.path.join(BASE, "Cronograma Mto Preventivo 3 _ Equipos 1.xlsx")
DATA = os.path.join(BASE, "Data_PCTriage.xlsx")
OUT = os.path.join(BASE, "Cumplimiento_Cronograma_MT3.xlsx")

FECHA_REF = hoy()
CATS = ["Equipo de escritorio", "Equipo portátil", "Escáner", "Impresora",
        "Lector biométrico", "Lector de banda pin pad", "Lector de código",
        "Monitor", "PAD de firmas", "Servidor", "Tablet"]

NAVY = "1F4E78"


def C(h):
    return "FF" + h


FILL = {
    "CUMPLIDA": ("E3F2E8", "1E7A43"),
    "VENCIDA CON PENDIENTES (Fact Si)": ("FBE6E5", "B03A34"),
    "EN EJECUCIÓN": ("E4EEF8", "1F4E78"),
    "ATRASADA SIN INICIAR (Fact Si)": ("FBE6E5", "B03A34"),
    "ADELANTADA": ("FDF0DC", "9A6410"),
    "PROGRAMADA": ("ECEEF1", "5A6572"),
    "REPROGRAMADA": ("D9D2E9", "5B2C8F"),
}


def cargar_cronograma():
    c = pd.read_excel(CRONO, sheet_name="Hoja1")
    c = c[c["tipo"].isin(["Oficina", "DG"])].copy()
    c["Ini"] = pd.to_datetime(c["Fecha Inicio"], errors="coerce")
    c["Fin"] = pd.to_datetime(c["Fecha Fin"], errors="coerce")
    c["SBAN"] = pd.to_numeric(c["SBAN"], errors="coerce")
    c["Crono"] = c[CATS].apply(pd.to_numeric, errors="coerce").sum(axis=1)
    if "ESTADO" not in c.columns:
        c["ESTADO"] = "Programada"
    g = c.groupby("SBAN", dropna=False).agg(
        Nombre=("Nombre Oficina", lambda s: " / ".join(dict.fromkeys(str(x).strip() for x in s if pd.notna(x) and str(x).strip()))),
        Regional=("Jefaturas  Operaciones Regional", lambda s: s.mode().iloc[0] if len(s) else ""),
        Estado=("ESTADO", lambda s: " / ".join(sorted(set(str(x) for x in s if pd.notna(x)))) if s.nunique() > 1 else s.mode().iloc[0]),
        Ini=("Ini", "min"), Fin=("Fin", "max"), Crono=("Crono", "sum")).reset_index()
    return g


def cargar_data():
    d = pd.read_excel(DATA, sheet_name="Hoja1")
    d["SBAN"] = pd.to_numeric(d["SBAN"], errors="coerce")
    d["mt"] = d["Consecutivo mantenimiento 3"].astype(str).str.upper().str.startswith("MT")
    d["si"] = d["Facturable"].astype(str).str.strip() == "Si"
    d["no"] = d["Facturable"].astype(str).str.strip() == "No"
    d["H"] = pd.to_datetime(d["Fecha de mantenimiento 3"], errors="coerce")
    d["mtSi"] = d["si"] & d["mt"]
    d["mtNo"] = d["no"] & d["mt"]
    g = d.groupby("SBAN").agg(
        Ele=("Serial", "size"), MT=("mt", "sum"),
        EleSi=("si", "sum"), MTSi=("mtSi", "sum"),
        EleNo=("no", "sum"), MTNo=("mtNo", "sum"),
        MinMT=("H", "min"), MaxMT=("H", "max")).reset_index()
    return d, g


def cargar_novedades():
    try:
        n = pd.read_excel(CRONO, sheet_name="NOVEDADES ACTIVOS")
        n["SBAN"] = pd.to_numeric(n["SBAN"], errors="coerce")
        tmp = n.groupby("SBAN")["ESTADO"].apply(
            lambda s: ", ".join(f"{k} ({v})" for k, v in s.value_counts().items()))
        return tmp.to_dict(), n
    except Exception:  # noqa: BLE001
        return {}, pd.DataFrame()


def estado_por_fechas(ini, fin):
    if pd.isna(ini) or pd.isna(fin):
        return "SIN FECHA"
    if fin < FECHA_REF:
        return "Cerrada"
    if ini <= FECHA_REF:
        return "En proceso"
    return "Programada"


def main():
    cg = cargar_cronograma()
    data_raw, da = cargar_data()
    nov_dict, nov_df = cargar_novedades()

    m = cg.merge(da, on="SBAN", how="left")
    for col in ["Ele", "MT", "EleSi", "MTSi", "EleNo", "MTNo"]:
        m[col] = m[col].fillna(0).astype(int)
    m["Crono"] = m["Crono"].fillna(0).astype(int)
    m["PendSi"] = (m["EleSi"] - m["MTSi"]).clip(lower=0)
    m["PendNo"] = (m["EleNo"] - m["MTNo"]).clip(lower=0)
    m["PendInv"] = (m["Ele"] - m["MT"]).clip(lower=0)
    m["AvanceSi"] = (m["MTSi"] / m["EleSi"]).where(m["EleSi"] > 0, None)
    m["EstadoFecha"] = m.apply(lambda r: estado_por_fechas(r["Ini"], r["Fin"]), axis=1)
    m["Dias plan"] = ((m["Fin"] - m["Ini"]).dt.days + 1).fillna(0).astype(int)

    def cumplimiento(r):
        e = r["Estado"]
        if e == "Finalizada":
            if r["PendSi"] == 0 and r["EleSi"] > 0:
                return "CUMPLIDA"
            if r["EleSi"] == 0:
                return "CUMPLIDA"
            return "VENCIDA CON PENDIENTES (Fact Si)"
        if e == "En proceso":
            if r["EleSi"] > 0 and r["PendSi"] == 0:
                return "EN EJECUCIÓN"
            if r["MTSi"] > 0:
                return "EN EJECUCIÓN"
            return "ATRASADA SIN INICIAR (Fact Si)"
        if e == "Programada":
            return "ADELANTADA" if r["MTSi"] > 0 else "PROGRAMADA"
        if e == "Reprogramada":
            return "REPROGRAMADA"
        return e if str(e).strip() else "SIN ESTADO"

    m["Cumplimiento"] = m.apply(cumplimiento, axis=1)

    def dias_desv(r):
        c = r["Cumplimiento"]
        if c == "VENCIDA CON PENDIENTES (Fact Si)":
            dias = (FECHA_REF - r["Fin"]).days if pd.notna(r["Fin"]) else 0
            return max(0, dias)
        if c == "ATRASADA SIN INICIAR (Fact Si)":
            fin = r["Fin"] if pd.notna(r["Fin"]) else r["Ini"]
            return max(0, (FECHA_REF - fin).days)
        if c == "ADELANTADA":
            return max(0, (r["Ini"] - FECHA_REF).days)
        return 0

    m["Dias desv"] = m.apply(dias_desv, axis=1)

    def coherencia(r):
        e, ef = r["Estado"], r["EstadoFecha"]
        if e == "Finalizada" and ef in ("En proceso", "Programada"):
            return "Revisar: Finalizada con fecha fin futura"
        if e in ("En proceso", "Programada") and ef == "Cerrada":
            return "Revisar: fecha fin vencida sin marcar Finalizada"
        if e == "Programada" and ef == "En proceso":
            return "Revisar: programada con fecha de inicio ya pasada"
        return "Coherente"

    m["Coherencia"] = m.apply(coherencia, axis=1)

    def observacion(r):
        o = []
        if r["PendSi"]:
            o.append(f"Fact Si pendientes: {r['PendSi']} de {r['EleSi']}")
        if r["PendNo"]:
            o.append(f"No fact pendientes (COLSOF): {r['PendNo']} de {r['EleNo']}")
        if r["Cumplimiento"] == "ADELANTADA":
            o.append("MT3 Fact Si antes de la fecha programada")
        if r["Coherencia"] != "Coherente":
            o.append(r["Coherencia"])
        nv = nov_dict.get(r["SBAN"])
        if nv:
            o.append("Novedad activa: " + nv)
        return " | ".join(o) if o else "Conforme"

    m["Observación"] = m.apply(observacion, axis=1)
    m["SBAN_TXT"] = m["SBAN"].astype(int).astype(str).str.zfill(5)

    # ---------------------------------------- escritura hoja control
    cols = ["Regional", "SBAN_TXT", "Nombre", "Estado", "EstadoFecha",
            "Cumplimiento", "Ini", "Fin", "Dias plan", "Crono",
            "EleSi", "MTSi", "PendSi", "AvanceSi",
            "EleNo", "MTNo", "PendNo", "Ele", "MT", "PendInv",
            "Dias desv", "Coherencia", "MinMT", "MaxMT", "Observación"]
    renom = {"SBAN_TXT": "SBAN", "Nombre": "Oficina", "EstadoFecha": "Estado (fechas)",
             "Ini": "Fecha Inicio", "Fin": "Fecha Fin", "Crono": "Equipos programados",
             "EleSi": "Fact Si - Total", "MTSi": "Fact Si - MT3",
             "PendSi": "Fact Si - Pend.", "AvanceSi": "Progreso Fact Si %",
             "EleNo": "No fact - Total", "MTNo": "No fact - MT3",
             "PendNo": "No fact - Pend.", "Ele": "Inventario total", "MT": "MT total",
             "PendInv": "Pend. inventario", "Dias desv": "Días desv.",
             "MinMT": "1er MT3 real", "MaxMT": "Último MT3 real"}
    ctl = m.sort_values(["Estado", "Cumplimiento", "Regional", "Nombre"])[cols].rename(columns=renom)

    with pd.ExcelWriter(OUT, engine="openpyxl") as w:
        ctl.to_excel(w, sheet_name="CONTROL POR OFICINA", index=False)
        sban_crono = set(cg["SBAN"].dropna().unique())
        sin = da[~da["SBAN"].isin(sban_crono)].copy()
        sin["SBAN_TXT"] = sin["SBAN"].astype(int).astype(str).str.zfill(5)
        nom = data_raw.groupby("SBAN")["Oficina"].agg(lambda s: s.mode().iloc[0] if len(s) else "")
        sin = sin.merge(nom.rename("Oficina"), on="SBAN", how="left")
        sin_out = pd.DataFrame({
            "SBAN": sin["SBAN_TXT"], "Oficina": sin["Oficina"],
            "Fact Si Total": sin["EleSi"], "Fact Si MT3": sin["MTSi"],
            "No fact Total": sin["EleNo"], "No fact MT3": sin["MTNo"],
            "Total inventario": sin["Ele"], "MT total": sin["MT"]})
        sin_out.to_excel(w, sheet_name="SIN PROGRAMAR", index=False)
        if not nov_df.empty:
            nov_df.to_excel(w, sheet_name="NOVEDADES ACTIVOS", index=False)

    wb = load_workbook(OUT)
    ws = wb.create_sheet("RESUMEN", 0)
    escribir_resumen(ws, m)
    estilo_control(wb["CONTROL POR OFICINA"])
    for hoja in ["SIN PROGRAMAR", "NOVEDADES ACTIVOS"]:
        if hoja in wb.sheetnames:
            h = wb[hoja]
            for cell in h[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor=C(NAVY))
            h.freeze_panes = "A2"
            h.auto_filter.ref = h.dimensions
    wb.save(OUT)

    # -------------------------------------------- resumen consola
    print("Archivo cronograma:", os.path.basename(CRONO), "| Fecha referencia:", FECHA_REF.date())
    print("Oficinas:", len(m), "| Fact Si total:", int(m['EleSi'].sum()),
          "| Fact Si MT3:", int(m['MTSi'].sum()),
          "| PROGRESO Fact Si:", f"{m['MTSi'].sum()/m['EleSi'].sum()*100:.1f}%")
    print("No fact total:", int(m['EleNo'].sum()), "| No fact MT3:", int(m['MTNo'].sum()))
    print("Estado oficinas:", m['Estado'].value_counts().to_dict())
    print("Cumplimiento:", m['Cumplimiento'].value_counts().to_dict())
    print("Desviaciones -> Vencidas Si:", int((m['Cumplimiento'] == 'VENCIDA CON PENDIENTES (Fact Si)').sum()),
          "| Atrasadas Si:", int((m['Cumplimiento'] == 'ATRASADA SIN INICIAR (Fact Si)').sum()),
          "| Adelantadas:", int((m['Cumplimiento'] == 'ADELANTADA').sum()),
          "| Reprogramadas:", int(m['Estado'].eq('Reprogramada').sum()))
    inc = m[m['Coherencia'] != 'Coherente']
    if len(inc):
        print("Incoherencias estado-vs-fechas:", len(inc))
        print(inc[['SBAN_TXT', 'Nombre', 'Estado', 'EstadoFecha']].head(10).to_string(index=False))
    print("Archivo salida:", OUT)


def escribir_resumen(ws, m):
    ws.sheet_view.showGridLines = False
    thin = Side(style="thin", color=C("BFBFBF"))
    bd = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws["A1"] = "CUMPLIMIENTO / DESVIACIÓN CRONOGRAMA MT3 - PROGRESO SOBRE SI FACTURABLE"
    ws["A1"].font = Font(size=15, bold=True, color=C(NAVY))
    ws.merge_cells("A1:P1")
    ws["A2"] = ("Estado de oficina = columna ESTADO del cronograma. El avance (%) se mide "
                "con Facturable = Si; los No facturables (COLSOF) se informan aparte.")
    ws["A2"].font = Font(size=9, italic=True, color=C("595959"))
    ws.merge_cells("A2:P2")

    tsi = int(m["EleSi"].sum()); mtsi = int(m["MTSi"].sum()); psi = int(m["PendSi"].sum())
    tno = int(m["EleNo"].sum()); mtno = int(m["MTNo"].sum()); pno = int(m["PendNo"].sum())

    def cnt_estado(serie, palabra):
        return sum(1 for x in serie
                   for t in str(x).split("/") if t.strip() == palabra)

    kpis = [
        ("Fecha de referencia", FECHA_REF.date()),
        ("Oficinas en cronograma", len(m)),
        ("", ""),
        ("ESTADO: Programadas", cnt_estado(m["Estado"], "Programada")),
        ("ESTADO: En proceso", cnt_estado(m["Estado"], "En proceso")),
        ("ESTADO: Finalizadas", cnt_estado(m["Estado"], "Finalizada")),
        ("ESTADO: Reprogramadas", cnt_estado(m["Estado"], "Reprogramada")),
        ("", ""),
        ("FACTURABLES SI - Total elementos", tsi),
        ("FACTURABLES SI - MT3 realizados", mtsi),
        ("FACTURABLES SI - Pendientes", psi),
        ("PROGRESO FACTURABLE SI (%)", f"{mtsi/tsi*100:.1f}%".replace(".", ",") if tsi else "n/d"),
        ("", ""),
        ("NO FACTURABLES - Total (COLSOF)", tno),
        ("NO FACTURABLES - MT3 realizados", mtno),
        ("NO FACTURABLES - Pendientes", pno),
        ("", ""),
        ("CUMPLIDAS", int(m['Cumplimiento'].eq('CUMPLIDA').sum())),
        ("VENCIDAS con pend. Fact Si (revisar)", int(m['Cumplimiento'].eq('VENCIDA CON PENDIENTES (Fact Si)').sum())),
        ("ATRASADAS sin iniciar Fact Si (revisar)", int(m['Cumplimiento'].eq('ATRASADA SIN INICIAR (Fact Si)').sum())),
        ("ADELANTADAS (revisar)", int(m['Cumplimiento'].eq('ADELANTADA').sum())),
        ("INCOHERENCIAS estado vs fechas", int((m['Coherencia'] != 'Coherente').sum())),
    ]
    r0 = 4
    for j, h in enumerate(["INDICADOR", "VALOR"], 1):
        c = ws.cell(r0, j, h)
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=C(NAVY)); c.border = bd
    for i, (k, v) in enumerate(kpis):
        fila = r0 + 1 + i
        a = ws.cell(fila, 1, k); b = ws.cell(fila, 2, v)
        a.border = bd; b.border = bd
        if isinstance(v, int):
            b.number_format = "#,##0"; b.font = Font(bold=True, color=C(NAVY))

    grp = m.groupby("Regional").agg(
        Oficinas=("Nombre", "size"),
        Programadas=("Estado", lambda s: cnt_estado(s, "Programada")),
        En_proceso=("Estado", lambda s: cnt_estado(s, "En proceso")),
        Finalizadas=("Estado", lambda s: cnt_estado(s, "Finalizada")),
        Reprogramadas=("Estado", lambda s: cnt_estado(s, "Reprogramada")),
        EleSi=("EleSi", "sum"), MTSi=("MTSi", "sum"),
        EleNo=("EleNo", "sum"), MTNo=("MTNo", "sum"),
        Vencidas=("Cumplimiento", lambda s: int((s == 'VENCIDA CON PENDIENTES (Fact Si)').sum())),
        Atrasadas=("Cumplimiento", lambda s: int((s == 'ATRASADA SIN INICIAR (Fact Si)').sum())),
        Adelantadas=("Cumplimiento", lambda s: int((s == 'ADELANTADA').sum())),
    ).reset_index()
    grp["PendSi"] = grp["EleSi"] - grp["MTSi"]
    grp["ProgresoSi"] = (grp["MTSi"] / grp["EleSi"] * 100).round(1)
    tot = pd.DataFrame([{
        "Regional": "TOTAL", "Oficinas": grp["Oficinas"].sum(),
        "Programadas": grp["Programadas"].sum(), "En_proceso": grp["En_proceso"].sum(),
        "Finalizadas": grp["Finalizadas"].sum(), "Reprogramadas": grp["Reprogramadas"].sum(),
        "EleSi": grp["EleSi"].sum(), "MTSi": grp["MTSi"].sum(),
        "EleNo": grp["EleNo"].sum(), "MTNo": grp["MTNo"].sum(),
        "Vencidas": grp["Vencidas"].sum(), "Atrasadas": grp["Atrasadas"].sum(),
        "Adelantadas": grp["Adelantadas"].sum(), "PendSi": grp["PendSi"].sum(),
        "ProgresoSi": round(grp["MTSi"].sum() / grp["EleSi"].sum() * 100, 1)}])
    grp = pd.concat([grp, tot], ignore_index=True)
    cab = ["Regional", "Oficinas", "Prog.", "En proceso", "Final.", "Reprog.",
           "Fact Si Total", "Fact Si MT3", "Fact Si Pend", "Progreso Si %",
           "No fact Total", "No fact MT3", "Vencidas Si", "Atrasadas Si", "Adelantadas"]
    fila = r0 + 1 + len(kpis) + 2
    ws.cell(fila, 1, "RESUMEN POR REGIONAL (progreso = Facturable Si)").font = Font(size=12, bold=True, color=C(NAVY))
    fila += 1
    for j, h in enumerate(cab, 1):
        c = ws.cell(fila, j, h)
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=C(NAVY)); c.border = bd
    for _, r in grp.iterrows():
        fila += 1
        vals = [r["Regional"], r["Oficinas"], r["Programadas"], r["En_proceso"],
                r["Finalizadas"], r["Reprogramadas"], r["EleSi"], r["MTSi"],
                r["PendSi"], f"{r['ProgresoSi']:.1f}%".replace(".", ","),
                r["EleNo"], r["MTNo"], r["Vencidas"], r["Atrasadas"], r["Adelantadas"]]
        for j, v in enumerate(vals, 1):
            c = ws.cell(fila, j, v); c.border = bd
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                c.number_format = "#,##0"
        if r["Regional"] == "TOTAL":
            for j in range(1, len(vals) + 1):
                ws.cell(fila, j).font = Font(bold=True)
                ws.cell(fila, j).fill = PatternFill("solid", fgColor=C("E8EEF4"))
    for col, w in zip("ABCDEFGHIJKLMNO", [22, 10, 10, 12, 11, 11, 12, 12, 12, 12, 12, 11, 11, 11, 11]):
        ws.column_dimensions[col].width = w


def estilo_control(ws):
    col_w = {"Regional": 15, "SBAN": 9, "Oficina": 36, "Estado": 14, "Estado (fechas)": 13,
             "Cumplimiento": 24, "Fecha Inicio": 12, "Fecha Fin": 12, "Días plan": 9,
             "Equipos programados": 12, "Fact Si - Total": 12, "Fact Si - MT3": 12,
             "Fact Si - Pend.": 12, "Progreso Fact Si %": 15, "No fact - Total": 12,
             "No fact - MT3": 12, "No fact - Pend.": 12, "Inventario total": 12,
             "MT total": 10, "Pend. inventario": 13, "Días desv.": 10, "Coherencia": 30,
             "1er MT3 real": 12, "Último MT3 real": 14, "Observación": 62}
    hdrs = [c.value for c in ws[1]]
    date_idx = {hdrs.index(n) for n in ("Fecha Inicio", "Fecha Fin", "1er MT3 real", "Último MT3 real") if n in hdrs}
    for i, h in enumerate(hdrs, 1):
        cell = ws.cell(1, i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=C(NAVY))
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = col_w.get(h, 12)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    cumple_idx = hdrs.index("Cumplimiento")
    for r in ws.iter_rows(min_row=2):
        cumple = r[cumple_idx].value
        if cumple in FILL:
            f, t = FILL[cumple]
            for c in r:
                c.fill = PatternFill("solid", fgColor=C(f))
                if c.value is not None and str(c.value).strip() != "":
                    c.font = Font(color=C(t))
        for idx in date_idx:
            if idx < len(r) and hasattr(r[idx].value, "strftime"):
                r[idx].number_format = "yyyy-mm-dd"
        av = r[hdrs.index("Progreso Fact Si %")] if "Progreso Fact Si %" in hdrs else None
        if av is not None and isinstance(av.value, (int, float)) and not isinstance(av.value, bool):
            av.number_format = "0.0%"


if __name__ == "__main__":
    main()
