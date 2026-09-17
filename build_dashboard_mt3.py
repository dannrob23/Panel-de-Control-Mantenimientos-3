# -*- coding: utf-8 -*-
"""
Dashboard Seguimiento Novedades Inventario - Mantenimiento 3 (MT3)
Fuente : Data_PCTriage.xlsx (hoja Hoja1). 1 fila = 1 elemento (Serial unico).
Salida : Dashboard_Seguimiento_MT3.xlsx
Reglas : AN='Si' -> BANCO | AN='No' -> COLSOF
         Diferencia = Elementos - Con MT3 (faltan por mantenimiento 3)
"""
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, PieChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import DataBarRule

SRC = r"C:\Users\darobles\Documents\PROYECTOS\CONTROL INVENTARIO MANTENIMIENTOS 3\Data_PCTriage.xlsx"
OUT = r"C:\Users\darobles\Documents\PROYECTOS\CONTROL INVENTARIO MANTENIMIENTOS 3\Dashboard_Seguimiento_MT3.xlsx"

NAVY = "1F4E78"; BLUE = "2E75B6"; ORANGE = "ED7D31"
GREEN = "70AD47"; RED = "C00000"; LIGHT = "D9E1F2"; GRAY = "595959"
def C(h): return "FF" + h

df = pd.read_excel(SRC, sheet_name="Hoja1")
df["Serial"] = df["Serial"].astype(str)
df["G"] = df["Consecutivo mantenimiento 3"]
df["H"] = df["Fecha de mantenimiento 3"]
df["tieneMT3"] = df["G"].notna()
df["Regional"] = df["Regional"].fillna("SIN REGIONAL")
df["Oficina"] = df["Oficina"].fillna("SIN NOMBRE")
for c in ["Fecha de asignación", "Fecha de mantenimiento 1", "Fecha de mantenimiento 2"]:
    df[c] = pd.to_datetime(df[c], errors="coerce")
df["H"] = pd.to_datetime(df["H"], errors="coerce")

AN_SI = df["Facturable"].eq("Si").sum()
AN_NO = df["Facturable"].eq("No").sum()
assert AN_SI + AN_NO == len(df)

# Tickets MT3 que aparecen en mas de un SBAN (posibles duplicados / sobrantes)
tk_sban = df[df["G"].notna()].groupby("G")["SBAN"].nunique()
shared = set(tk_sban[tk_sban > 1].index)
sban_por_tk = df[df["G"].isin(shared)].groupby("G")["SBAN"].apply(lambda s: sorted(set(s)))

grp = df.groupby("SBAN")
ser_elementos = grp.size()
res = pd.DataFrame({
    "SBAN": ser_elementos.index,
    "Oficina": grp["Oficina"].agg(lambda s: s.mode().iloc[0]).values,
    "Regional": grp["Regional"].agg(lambda s: s.mode().iloc[0]).values,
    "Elementos": ser_elementos.values,
    "Con MT3": grp["tieneMT3"].sum().values,
    "BANCO Si": grp["Facturable"].apply(lambda s: (s == "Si").sum()).values,
    "COLSOF No": grp["Facturable"].apply(lambda s: (s == "No").sum()).values,
    "Tickets MT3": grp["G"].apply(lambda s: s.dropna().nunique()).values,
})
valfac = df[df["Facturable"] == "Si"].groupby("SBAN")["Valor facturación"].sum()
res["Valor facturacion"] = res["SBAN"].map(valfac).fillna(0.0).values
res["Pendientes (Diferencia)"] = res["Elementos"] - res["Con MT3"]
res["Avance %"] = res["Con MT3"] / res["Elementos"]

def observacion(row):
    o = []
    if row["Pendientes (Diferencia)"] > 0:
        o.append("FALTAN %d por MT3" % row["Pendientes (Diferencia)"])
    elif row["Con MT3"] > 0:
        o.append("Completo MT3")
    if row["Oficina"] == "SIN NOMBRE":
        o.append("Oficina sin nombre")
    if row["SBAN"] in df.groupby("SBAN")["Oficina"].nunique().to_dict() and \
       df.groupby("SBAN")["Oficina"].nunique()[row["SBAN"]] > 1:
        o.append("SBAN con varios nombres de oficina")
    tickets = df[df["SBAN"] == row["SBAN"]]["G"].dropna()
    sh = [t for t in tickets if t in shared]
    if sh:
        o.append("SOBRA? ticket(s) %s en otros SBAN" % ",".join(str(t) for t in sorted(set(sh))[:3]))
    return "; ".join(o) if o else "OK"

res["Observación"] = res.apply(observacion, axis=1)
res = res.sort_values(["Regional", "Pendientes (Diferencia)"], ascending=[True, False]).reset_index(drop=True)

# ---- Totales de control ----
TOT = res.sum(numeric_only=True)
totales = {
    "Elementos": int(TOT["Elementos"]), "Con MT3": int(TOT["Con MT3"]),
    "Pendientes": int(TOT["Pendientes (Diferencia)"]),
    "BANCO": int(TOT["BANCO Si"]), "COLSOF": int(TOT["COLSOF No"]),
    "Tickets": int(df["G"].dropna().nunique()),
    "Valor": float(TOT["Valor facturacion"]),
}

# ---- Datos para graficas ----
reg_order = ["Antioquia", "Bogotá", "Cafetera", "Costa", "Dirección General",
             "Occidente", "Oriente", "Santanderes", "Sur", "SIN REGIONAL"]
regs = [r for r in reg_order if r in set(res["Regional"])]
regav = res.groupby("Regional", sort=False).agg(
    Elementos=("Elementos", "sum"), ConMT3=("Con MT3", "sum"),
    Pend=("Pendientes (Diferencia)", "sum"),
    BANCO=("BANCO Si", "sum"), COLSOF=("COLSOF No", "sum")).reindex(regs)
regav["Avance %"] = regav["ConMT3"] / regav["Elementos"]

top = res.nlargest(15, "Pendientes (Diferencia)").copy()
top["Etiqueta"] = top["SBAN"].astype(str).str.zfill(5) + " " + top["Oficina"].str[:20]

dia = df[df["H"].notna()].groupby(df["H"].dt.date)["Serial"].count().sort_index()

# ================= ESCRIBIR LIBRO =================
with pd.ExcelWriter(OUT, engine="openpyxl") as w:
    res[["Regional", "SBAN", "Oficina", "Elementos", "Con MT3",
         "Pendientes (Diferencia)", "Avance %", "BANCO Si", "COLSOF No",
         "Tickets MT3", "Valor facturacion", "Observación"]].to_excel(
        w, sheet_name="CONTROL POR OFICINA", index=False)
    falt = df[~df["tieneMT3"]][["Regional", "SBAN", "Oficina", "Tipo ubicación", "Categoría",
                                 "Descripción", "Modelo", "Serial", "Placa", "Estado",
                                 "Estado interno", "Facturable"]]
    falt.to_excel(w, sheet_name="FALTANTES MT3", index=False)

wb = load_workbook(OUT)
ws_c = wb["CONTROL POR OFICINA"]
ws_f = wb["FALTANTES MT3"]

# ---- Hoja DATOS (auxiliar para graficas, oculta) ----
ws_d = wb.create_sheet("DATOS GRAFICOS")
ws_d.sheet_state = "hidden"
ws_d["A1"] = "Regional"; ws_d["B1"] = "Con MT3"; ws_d["C1"] = "Pendientes"; ws_d["D1"] = "Elementos"
for i, (reg, r) in enumerate(regav.iterrows(), start=2):
    ws_d.cell(i, 1, reg); ws_d.cell(i, 2, int(r["ConMT3"]))
    ws_d.cell(i, 3, int(r["Pend"])); ws_d.cell(i, 4, int(r["Elementos"]))
nreg = len(regav)
ws_d["A14"] = "Tipo"; ws_d["B14"] = "Elementos"
ws_d["A15"] = "BANCO (Facturable Si)"; ws_d["B15"] = AN_SI
ws_d["A16"] = "COLSOF (No Facturable)"; ws_d["B16"] = AN_NO
ws_d["A19"] = "Fecha"; ws_d["B19"] = "MT3 realizados"
for i, (d, n) in enumerate(dia.items(), start=20):
    ws_d.cell(i, 1, d.strftime("%d-%m-%Y")); ws_d.cell(i, 2, int(n))
ndia = len(dia)
ws_d["A30"] = "SBAN - Oficina"; ws_d["B30"] = "Pendientes"
for i, (_, r) in enumerate(top.iterrows(), start=31):
    ws_d.cell(i, 1, r["Etiqueta"]); ws_d.cell(i, 2, int(r["Pendientes (Diferencia)"]))
ntop = len(top)

# ================= DASHBOARD =================
ws = wb.create_sheet("DASHBOARD", 0)
thin = Side(style="thin", color=C("BFBFBF"))
bord = Border(left=thin, right=thin, top=thin, bottom=thin)
ws.sheet_view.showGridLines = False

ws.merge_cells("A1:L2")
ws["A1"] = "DASHBOARD SEGUIMIENTO NOVEDADES INVENTARIO  -  MANTENIMIENTO 3 (MT3)"
ws["A1"].font = Font(size=18, bold=True, color=C(NAVY))
ws.merge_cells("A3:L3")
ws["A3"] = ("Fuente: Data_PCTriage.xlsx (Hoja1) | 1 fila = 1 elemento | "
            "AN Si = BANCO / AN No = COLSOF | Diferencia = Elementos - Con MT3")
ws["A3"].font = Font(size=9, italic=True, color=C(GRAY))

def card(row0, col, value, label, fill):
    v = ws.cell(row0, col, value); v.font = Font(size=20, bold=True, color=C(NAVY))
    v.fill = PatternFill("solid", fgColor=C(LIGHT)); v.alignment = Alignment(horizontal="center")
    l = ws.cell(row0 + 1, col, label)
    l.font = Font(size=8, color=C(GRAY)); l.alignment = Alignment(horizontal="center", wrap_text=True)
    l.fill = PatternFill("solid", fgColor=C("F2F2F2"))
    for r in (row0, row0 + 1):
        ws.cell(r, col).border = bord
    ws.column_dimensions[get_column_letter(col)].width = 13

R = 5
card(R, 1, f"{totales['Elementos']:,}".replace(",", "."), "TOTAL ELEMENTOS", LIGHT)
card(R, 2, f"{totales['Con MT3']:,}".replace(",", "."), "CON MT3 (col G)", LIGHT)
card(R, 3, f"{totales['Con MT3']/totales['Elementos']:.1%}".replace(",", "."), "AVANCE %", LIGHT)
card(R, 4, f"{totales['Pendientes']:,}".replace(",", "."), "PENDIENTES (faltan MT3)", "FCE4D6")
card(R, 5, f"{totales['Tickets']:,}".replace(",", "."), "TICKETS MT3 UNICOS", LIGHT)
card(R, 6, f"{totales['BANCO']:,}".replace(",", "."), "BANCO (Facturable Si)", LIGHT)
card(R, 7, f"{totales['COLSOF']:,}".replace(",", "."), "COLSOF (No Facturable)", LIGHT)
card(R, 8, f"{totales['Valor']:,.0f}".replace(",", "."), "VALOR FACTURACION $", LIGHT)

ch1 = BarChart(); ch1.type = "col"; ch1.grouping = "stacked"; ch1.overlap = 100
ch1.title = "Realizado vs Pendiente por Regional"
d1 = Reference(ws_d, min_col=2, min_row=1, max_col=3, max_row=nreg + 1)
cats = Reference(ws_d, min_col=1, min_row=2, max_row=nreg + 1)
ch1.add_data(d1, titles_from_data=True); ch1.set_categories(cats)
ch1.series[0].graphicalProperties.solidFill = C(GREEN)
ch1.series[1].graphicalProperties.solidFill = C(RED)
ch1.width = 15; ch1.height = 8.5
ws.add_chart(ch1, "A10")

ch2 = PieChart(); ch2.title = "BANCO vs COLSOF"
d2 = Reference(ws_d, min_col=1, min_row=14, max_col=2, max_row=16)
ch2.add_data(d2, titles_from_data=True, from_rows=True)
ch2.set_categories(Reference(ws_d, min_col=1, min_row=15, max_row=16))
ch2.series[0].dataLabels = DataLabelList(); ch2.series[0].dataLabels.showPercent = True
ch2.width = 12; ch2.height = 8.5
ws.add_chart(ch2, "J10")

ch3 = BarChart(); ch3.type = "bar"
ch3.add_data(Reference(ws_d, min_col=2, min_row=30, max_row=30 + ntop), titles_from_data=True)
ch3.set_categories(Reference(ws_d, min_col=1, min_row=31, max_row=30 + ntop))
ch3.series[0].graphicalProperties.solidFill = C(RED)
ch3.width = 15; ch3.height = 9
ws.add_chart(ch3, "A30")

ch4 = LineChart()
ch4.add_data(Reference(ws_d, min_col=2, min_row=19, max_row=19 + ndia), titles_from_data=True)
ch4.set_categories(Reference(ws_d, min_col=1, min_row=20, max_row=19 + ndia))
ch4.series[0].graphicalProperties.line.solidFill = C(BLUE)
ch4.series[0].smooth = False
ch4.width = 12; ch4.height = 9
ws.add_chart(ch4, "J30")

nota = ws.cell(56, 1, "Nota: 1 fila = 1 elemento (Serial unico). La diferencia (Pendientes) nunca es negativa en esta base: "
                      "los casos 'SOBRA' se detectan cuando un mismo ticket MT3 (col G) aparece en mas de un SBAN y se marcan "
                      "en la columna Observacion de CONTROL POR OFICINA y en la hoja 'ALERTAS SOBRANTES'.")
nota.font = Font(size=9, italic=True, color=C(GRAY))
ws.merge_cells("A56:L58"); nota.alignment = Alignment(wrap_text=True, vertical="top")

# ================= CONTROL POR OFICINA =================
hdrs = [c.value for c in ws_c[1]]
widths = {"Regional": 17, "SBAN": 9, "Oficina": 26, "Elementos": 11, "Con MT3": 11,
          "Pendientes (Diferencia)": 13, "Avance %": 10, "BANCO Si": 11, "COLSOF No": 11,
          "Tickets MT3": 10, "Valor facturacion": 14, "Observación": 72}
for i, h in enumerate(hdrs, 1):
    L = get_column_letter(i)
    ws_c.column_dimensions[L].width = widths.get(h, 12)
    c = ws_c.cell(1, i); c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=C(NAVY))
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
for r in ws_c.iter_rows(min_row=2):
    r[0].alignment = Alignment(vertical="top")
    r[2].alignment = Alignment(vertical="top")
    r[3].number_format = "#,##0"; r[4].number_format = "#,##0"
    r[5].number_format = "#,##0"
    r[6].number_format = "0.0%"
    r[7].number_format = "#,##0"; r[8].number_format = "#,##0"; r[9].number_format = "#,##0"
    r[10].number_format = "#,##0"
    r[11].alignment = Alignment(vertical="top", wrap_text=True)
    if isinstance(r[1].value, (int, float)):
        r[1].number_format = "00000"
ws_c.freeze_panes = "A2"
n_datos = res.shape[0]
ws_c.auto_filter.ref = f"A1:L{n_datos + 1}"
last = n_datos + 2
ws_c.cell(last, 1, "TOTAL")
for i, h in enumerate(hdrs, 1):
    col = get_column_letter(i)
    v = ws_c.cell(last, i)
    v.font = Font(bold=True); v.fill = PatternFill("solid", fgColor=C(LIGHT))
    if i in (4, 5, 6, 8, 9, 10, 11):
        tot = {4: totales["Elementos"], 5: totales["Con MT3"], 6: totales["Pendientes"],
               8: totales["BANCO"], 9: totales["COLSOF"], 10: int(res["Tickets MT3"].sum()),
               11: totales["Valor"]}[i]
        v.value = tot; v.number_format = "#,##0"
    if i == 7:
        v.value = totales["Con MT3"] / totales["Elementos"]; v.number_format = "0.0%"
    if i == 12:
        v.value = "Ver KPIs en DASHBOARD"
ws_c.cell(last, 1).alignment = Alignment(horizontal="left")
cell_av = f"G2:G{n_datos + 1}"
ws_c.conditional_formatting.add(cell_av, DataBarRule(start_type="num", start_value=0,
    end_type="num", end_value=1, color=C(GREEN), showValue=True, minLength=None, maxLength=None))

# ================= FALTANTES MT3 =================
ws_f.freeze_panes = "A2"
ws_f.auto_filter.ref = ws_f.dimensions
for i, h in enumerate([c.value for c in ws_f[1]], 1):
    L = get_column_letter(i)
    c = ws_f.cell(1, i); c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=C(RED))
ws_f.column_dimensions["A"].width = 17; ws_f.column_dimensions["B"].width = 9
ws_f.column_dimensions["C"].width = 26; ws_f.column_dimensions["G"].width = 22
ws_f.column_dimensions["D"].width = 12; ws_f.column_dimensions["J"].width = 14
ws_f.column_dimensions["K"].width = 24
for r in ws_f.iter_rows(min_row=2, min_col=2, max_col=2):
    for c in r:
        if isinstance(c.value, (int, float)): c.number_format = "00000"
for r in ws_f.iter_rows(min_row=2, min_col=9, max_col=9):
    for c in r:
        if isinstance(c.value, (int, float)): c.number_format = "00000000"
ws_f.auto_filter.ref = f"A1:L{falt.shape[0] + 1}"

# ================= ALERTAS (sobrantes / inconsistencias) =================
ws_a = wb.create_sheet("ALERTAS SOBRANTES")
al = df[df["G"].isin(shared)]
ws_a.append(["Ticket MT3 (G)", "SBAN", "Oficina", "Serial", "Categoria", "Detalle"])
for _, r in al.iterrows():
    sbs = sban_por_tk[r["G"]]
    ws_a.append([r["G"], r["SBAN"], r["Oficina"], r["Serial"], r["Categoría"],
                 "Ticket en %d SBAN: %s" % (len(sbs), ",".join(str(s) for s in sbs))])
g_sin_h = df[df["G"].notna() & df["H"].isna()]
h_sin_g = df[df["G"].isna() & df["H"].notna()]
for n, txt in ((len(g_sin_h), "filas con consecutivo MT3 sin fecha"),
               (len(h_sin_g), "filas con fecha MT3 sin consecutivo")):
    if n:
        ws_a.append(["INCONSISTENCIA", "", "", "", "", "%d %s" % (n, txt)])
for i, h in enumerate([c.value for c in ws_a[1]], 1):
    c = ws_a.cell(1, i); c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=C(NAVY))
ws_a.freeze_panes = "A2"
ws_a.auto_filter.ref = f"A1:F{ws_a.max_row}"
ws_a.column_dimensions["A"].width = 16; ws_a.column_dimensions["B"].width = 9
ws_a.column_dimensions["C"].width = 26; ws_a.column_dimensions["D"].width = 15
ws_a.column_dimensions["E"].width = 12; ws_a.column_dimensions["F"].width = 60
for r in ws_a.iter_rows(min_row=2, min_col=2, max_col=2):
    for c in r:
        if isinstance(c.value, (int, float)): c.number_format = "00000"

wb.save(OUT)

# ================= VERIFICACION =================
print("OK DASHBOARD")
print("CONTROL filas:", n_datos, "| check suma Elementos:", int(res["Elementos"].sum()),
      "| Con MT3:", int(res["Con MT3"].sum()), "| Pend:", int(res["Pendientes (Diferencia)"].sum()))
print("BANCO:", totales["BANCO"], "| COLSOF:", totales["COLSOF"], "| Tickets:", totales["Tickets"],
      "| Valor $:", round(totales["Valor"]))
print("FALTANTES filas:", falt.shape[0], "| DASHBOARD sheets ok")
